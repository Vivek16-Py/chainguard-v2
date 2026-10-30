# -*- coding: utf-8 -*-
"""Build the ChainGuard v2 seminar report using the NMIET report template."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
REPORT = ROOT / "report"
FIG = REPORT / "diagrams"
OUT_DOCX = ROOT / "ChainGuard_v2_Technical_Seminar_Report.docx"

sys.path.insert(0, str(REPORT))
from content_a import BLOCKS_A  # noqa: E402
from content_b import BLOCKS_B  # noqa: E402

BLOCKS = BLOCKS_A + BLOCKS_B

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"

TITLE = (
    "CHAINGUARD V2: A TAMPER-EVIDENT LAST-MILE LOGISTICS CONTROL PLANE FOR "
    "SEALED DISPATCH, CUSTOMER-AUTHORISED CUSTODY BREAKS AND UNIT-EXACT "
    "SHRINKAGE RECONCILIATION"
)
SHORT_TITLE = "CHAINGUARD V2: A TAMPER-EVIDENT LAST-MILE LOGISTICS CONTROL PLANE"
TITLE_DISPLAY = (
    "ChainGuard v2: A Tamper-Evident Last-Mile Logistics Control Plane for Sealed "
    "Dispatch, Customer-Authorised Custody Breaks and Unit-Exact Shrinkage Reconciliation"
)
RUNNING_HEAD = (
    "CHAINGUARD V2: A TAMPER-EVIDENT LAST-MILE LOGISTICS CONTROL PLANE"
)
INSTITUTION = "NMIET, Computer Science and Engineering"

# template geometry (US Letter, see Seminar_Report_Edited_Header.pdf)
PAGE_W = 8.5
PAGE_H = 11.0
MARGIN_L = 1.45
MARGIN_R = 1.35
MARGIN_T = 1.0
MARGIN_B = 1.0
TEXT_W = PAGE_W - MARGIN_L - MARGIN_R

BODY_SIZE = 10.8
BODY_LINE = 1.14
CAPTION_SIZE = 10.4
TABLE_SIZE = 9.9

CHAPTER_TITLES = {
    "1": "1. Introduction",
    "2": "2. Literature Survey",
    "3": "3. Overview of the ChainGuard v2 Control Plane",
    "4": "4. Working / Methodology",
    "5": "5. Applications",
    "6": "6. Advantages and Limitations",
    "7": "7. Comparison with Existing Technologies",
    "8": "8. Future Scope",
    "9": "9. Conclusion",
}

_bookmark_id = [1000]


# ---------------------------------------------------------------------------
# low-level helpers
# ---------------------------------------------------------------------------
def set_rfonts(rpr, name):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), name)


def style_run(run, size=11.3, bold=None, italic=None, name=BODY_FONT, color=None, caps=False):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if caps:
        run.font.all_caps = True
    set_rfonts(run._element.get_or_add_rPr(), name)


TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_rich(paragraph, text, size=BODY_SIZE, italic=False, bold_all=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            style_run(run, size=size, bold=True, italic=italic or None)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            style_run(
                run,
                size=max(size - 1.6, 7.8),
                name=MONO_FONT,
                bold=bold_all or None,
                italic=italic or None,
            )
        else:
            run = paragraph.add_run(part)
            style_run(run, size=size, bold=bold_all or None, italic=italic or None)


CAP_RE = re.compile(r"^(Figure|Table)\s+(\d+\.\d+)[.:]\s*(.*?)(?:\.)?$")


def caption_parts(text):
    match = CAP_RE.match(text.strip())
    if not match:
        return text.strip(), text.strip()
    return f"{match.group(1)} {match.group(2)}", match.group(3).rstrip(".")


def caption_text(text):
    label, body = caption_parts(text)
    return f"{label}: {body}"


def add_bookmark(paragraph, name):
    _bookmark_id[0] += 1
    ident = str(_bookmark_id[0])
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), ident)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), ident)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_field(paragraph, instruction, placeholder="", size=BODY_SIZE, bold=None, tab_first=False):
    if tab_first:
        run = paragraph.add_run("\t")
        style_run(run, size=size)
    run = paragraph.add_run()
    style_run(run, size=size, bold=bold)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_el = OxmlElement("w:t")
    text_el.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text_el, end):
        run._r.append(element)
    return run


def right_tab(paragraph):
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(TEXT_W), alignment=WD_TAB_ALIGNMENT.RIGHT, leader=WD_TAB_LEADER.DOTS
    )


def para(
    text="",
    size=BODY_SIZE,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    before=0,
    after=3,
    line=BODY_LINE,
    left=0.0,
    first=0.0,
    keep=False,
    italic=False,
    bold_all=False,
):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if text.count("`") >= 6:
        align = WD_ALIGN_PARAGRAPH.LEFT
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if left:
        pf.left_indent = Inches(left)
    if first:
        pf.first_line_indent = Inches(first)
    pf.keep_with_next = keep
    if text:
        add_rich(p, text, size=size, italic=italic, bold_all=bold_all)
    return p


def center(text, size=BODY_SIZE, bold=False, before=0, after=0, italic=False, line=1.0, caps=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, italic=italic, caps=caps)
    return p


def spacer(points=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(points)
    pf.line_spacing = 1.0
    run = p.add_run("")
    run.font.size = Pt(2)
    return p


def add_item(marker, text, size=BODY_SIZE, left=0.36, hang=0.26, after=0.5):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(left)
    pf.first_line_indent = Inches(-hang)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = BODY_LINE
    pf.tab_stops.add_tab_stop(Inches(left))
    run = p.add_run(marker + "\t")
    style_run(run, size=size)
    add_rich(p, text, size=size)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_margins(table, left=0.06, right=0.06, top=0.02, bottom=0.02):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(value * 1440)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.insert_element_before(mar, "w:tblLook")


def no_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")


def bottom_rule(cell, size="4", color="808080"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    el = OxmlElement("w:bottom")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), size)
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)
    borders.append(el)
    tc_pr.append(borders)


def add_page_border(section):
    """Double rule around the page, matching the template's two-line frame."""
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:pgBorders"))
    if old is not None:
        sect_pr.remove(old)
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "thickThinMediumGap")
        element.set(qn("w:sz"), "12")
        element.set(qn("w:space"), "31" if edge in ("left", "right") else "24")
        element.set(qn("w:color"), "000000")
        pg_borders.append(element)
    sect_pr.insert_element_before(pg_borders, "w:lnNumType", "w:pgNumType", "w:cols", "w:titlePg")


def set_page_numbering(section, fmt=None, start=None):
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:pgNumType"))
    if old is not None:
        sect_pr.remove(old)
    if fmt is None and start is None:
        return
    pg = OxmlElement("w:pgNumType")
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    sect_pr.insert_element_before(
        pg, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
        "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings"
    )


def config_section(section):
    section.page_width = Inches(PAGE_W)
    section.page_height = Inches(PAGE_H)
    section.left_margin = Inches(MARGIN_L)
    section.right_margin = Inches(MARGIN_R)
    section.top_margin = Inches(MARGIN_T)
    section.bottom_margin = Inches(MARGIN_B)
    section.header_distance = Inches(0.10)
    section.footer_distance = Inches(0.16)
    add_page_border(section)


def add_page_number_footer(section, first_page=False):
    footer = section.first_page_footer if first_page else section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.clear()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.tab_stops.clear_all()
    pf.tab_stops.add_tab_stop(Inches(TEXT_W), alignment=WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(INSTITUTION)
    style_run(run, size=10)
    add_field(p, "PAGE", "1", size=10, tab_first=True)


def add_running_header(section, chapter_label=None, first_page=False):
    header = section.first_page_header if first_page else section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.clear()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.space_before = Pt(0)
    pf.space_after = Pt(14 if chapter_label else 0)
    pf.line_spacing = 1.0
    run = p.add_run(RUNNING_HEAD)
    style_run(run, size=9.5, bold=True, caps=True)
    if chapter_label:
        p2 = header.add_paragraph()
        pf = p2.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pf.right_indent = Inches(-0.62)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        run = p2.add_run(chapter_label)
        style_run(run, size=12, bold=True)


def blank_header(section):
    header = section.header
    header.is_linked_to_previous = False
    for p in header.paragraphs:
        p.clear()


def blank_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        p.clear()


def new_section():
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    config_section(section)
    return section


# ---------------------------------------------------------------------------
# document setup
# ---------------------------------------------------------------------------
doc = Document()

first = doc.sections[0]
config_section(first)
blank_header(first)
blank_footer(first)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(BODY_SIZE)
set_rfonts(normal.element.get_or_add_rPr(), BODY_FONT)
npf = normal.paragraph_format
npf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
npf.line_spacing = BODY_LINE
npf.space_before = Pt(0)
npf.space_after = Pt(3)


def config_heading(name, size, before, after, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    st = styles[name]
    st.font.name = BODY_FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.italic = italic
    st.font.color.rgb = RGBColor(0, 0, 0)
    rpr = st.element.get_or_add_rPr()
    set_rfonts(rpr, BODY_FONT)
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            if color.get(qn(attr)) is not None:
                del color.attrib[qn(attr)]
    pf = st.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.05
    pf.keep_with_next = True
    pf.keep_together = True
    pf.alignment = align


config_heading("Heading 1", 15, 0, 10)
config_heading("Heading 2", 13.2, 8, 3)
config_heading("Heading 3", 11.5, 6, 2, italic=True)

# Footer style in the default template carries a centre tab; replace it so the
# page number sits at the right margin like the reference layout.
try:
    footer_style = styles["Footer"]
    footer_style.paragraph_format.tab_stops.clear_all()
    footer_style.paragraph_format.tab_stops.add_tab_stop(
        Inches(TEXT_W), alignment=WD_TAB_ALIGNMENT.RIGHT
    )
except KeyError:
    pass

# TOC styles used by Word when the field is refreshed
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402

for toc_name, bold in (("TOC 1", True), ("TOC 2", False)):
    try:
        tst = styles[toc_name]
    except KeyError:
        tst = styles.add_style(toc_name, WD_STYLE_TYPE.PARAGRAPH)
    tst.font.name = BODY_FONT
    tst.font.size = Pt(11.3)
    tst.font.bold = bold
    tst.font.color.rgb = RGBColor(0, 0, 0)
    set_rfonts(tst.element.get_or_add_rPr(), BODY_FONT)
    tpf = tst.paragraph_format
    tpf.space_before = Pt(0)
    tpf.space_after = Pt(2)
    tpf.line_spacing = 1.05
    tpf.tab_stops.add_tab_stop(
        Inches(TEXT_W), alignment=WD_TAB_ALIGNMENT.RIGHT, leader=WD_TAB_LEADER.DOTS
    )

props = doc.core_properties
props.title = SHORT_TITLE
props.subject = "Technical Seminar Report"
props.author = "Vivek Rathod"
props.keywords = "ChainGuard, last-mile logistics, custody, shrinkage, reconciliation, React, TypeScript"
props.comments = "Technical seminar report, Computer Science and Engineering Department, NMIET."


def front_heading(text, bookmark=None, break_before=False, size=15, before=0, after=18):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    pf.page_break_before = break_before
    run = p.add_run(text)
    style_run(run, size=size, bold=True)
    if bookmark:
        add_bookmark(p, bookmark)
    return p


def rest_heading(text, before=22, after=12):
    return front_heading(text, size=15, before=before, after=after)


def letterhead():
    """Institute letterhead with the college and university logos."""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    no_borders(table)
    cell_margins(table, left=0.02, right=0.02, top=0.0, bottom=0.0)
    widths = (0.68, TEXT_W - 1.36, 0.68)
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)

    for index, filename, width in (
        (0, "logo_college.png", 0.62),
        (2, "logo_university.png", 0.58),
    ):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(FIG / filename), width=Inches(width))

    cell = table.rows[0].cells[1]
    cell.width = Inches(widths[1])
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    lines = (
        ("Nutan Maharashtra Vidya Prasarak Mandal’s (NMVPM’s)", 9.2),
        ("NUTAN MAHARASHTRA INSTITUTE OF ENGINEERING AND TECHNOLOGY (NMIET)", 12.4),
        ("An Autonomous Institute", 12.4),
        ("Under Administrative Support – Pimpri Chinchwad Education Trust (PCET)", 8.6),
    )
    for i, (text, size) in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0.5)
        pf.line_spacing = 1.0
        run = p.add_run(text)
        style_run(run, size=size, bold=True)

    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(2)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.tab_stops.add_tab_stop(Inches(TEXT_W / 2), alignment=WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Inches(TEXT_W), alignment=WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run("Approved by AICTE\tAccredited by NAAC\tAffiliated to SPPU")
    style_run(run, size=8.6, bold=True)
    spacer(10)


def student_table(rows, size=11.0):
    """Borderless name/PRN table used on the title and certificate pages."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    no_borders(table)
    cell_margins(table, left=0.05, right=0.05, top=0.01, bottom=0.01)
    widths = (2.55, 1.15)
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)

    for i, label in enumerate(("Student Name", "Roll No.")):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(label)
        style_run(run, size=size - 0.5, bold=True)
        bottom_rule(cell, size="4", color="000000")

    for name, prn in rows:
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, value in enumerate((name, prn)):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(value)
            style_run(run, size=size)
    spacer(6)


def signature_table(rows):
    """Guide / HOD / Principal signature block."""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    no_borders(table)
    cell_margins(table, left=0.02, right=0.02, top=0.0, bottom=0.0)
    width = TEXT_W / 3
    for i, (name, role, place) in enumerate(rows):
        table.columns[i].width = Inches(width)
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(name)
        style_run(run, 10.8)
        for text, bold in ((role, True), (place, True)):
            p2 = cell.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(1)
            p2.paragraph_format.space_after = Pt(0)
            p2.paragraph_format.line_spacing = 1.0
            run = p2.add_run(text)
            style_run(run, 10.2, bold=bold)


# ---------------------------------------------------------------------------
# title page
# ---------------------------------------------------------------------------
letterhead()
center("Technical Seminar Report on", 17, True, after=10)
center(TITLE, 15, True, after=8, line=1.12)
center("Project Reference: ChainGuard v2.4", 10.5, True, after=30)

center("SUBMITTED TO THE DEPARTMENT OF", 11.5, True, after=2)
center("COMPUTER SCIENCE & ENGINEERING", 11.5, True, after=36)

center("SUBMITTED BY", 13, True, after=14)
student_table((("Vivek Rathod", "B49"),))

center("Under The Guidance of", 13.5, True, before=44, after=4)
center("Miss. Sakshi Ambadkar", 13.5, True, after=106)

center("ACADEMIC YEAR : 2026–27", 13.5, True, after=10)
center(
    "Samarth Vidya Sankul, Vishnupuri, Talegaon Dabhade, Maharashtra 410507",
    8.8,
    after=0,
)

# ---------------------------------------------------------------------------
# front matter section: Roman numerals, footer only
# ---------------------------------------------------------------------------
front = new_section()
blank_header(front)
set_page_numbering(front, "upperRoman", 1)
add_page_number_footer(front)

# certificate
letterhead()
front_heading("CERTIFICATE", "bm_cert", size=18, after=16)
para(
    "This is to certify that the technical seminar report",
    11.5,
    WD_ALIGN_PARAGRAPH.CENTER,
    after=7,
)
para(
    "“" + TITLE_DISPLAY + "”",
    12,
    WD_ALIGN_PARAGRAPH.CENTER,
    after=7,
    line=1.15,
    bold_all=True,
)
para("Submitted by", 11.5, WD_ALIGN_PARAGRAPH.CENTER, after=5)
student_table((("Vivek Rathod", "B49"),), size=11)
para(
    "Has completed the technical seminar report satisfactorily in the subject of "
    "Technical Seminar (ELC342CSE) for the academic year 2026–2027 as prescribed in "
    "the syllabus of Savitribai Phule Pune University, Pune.",
    11,
    after=6,
)
para(
    "The work was carried out under the guidance of Miss. Sakshi Ambadkar and has "
    "not been submitted earlier to any other institute or university for the award "
    "of a degree.",
    11,
    after=0,
)

spacer(150)
signature_table(
    (
        ("Miss. Sakshi Ambadkar", "Seminar Guide", "Dept. of CSE"),
        ("Dr. Mahesh Wankhede", "HOD", "Dept. of CSE"),
        ("Dr. Pramod Patil", "Principal", "NMIET Pune"),
    )
)
spacer(22)
para("Place : Pune", 10.8, WD_ALIGN_PARAGRAPH.LEFT, after=1, line=1.1)
para("Date : ____________________", 10.8, WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.1)

# acknowledgement
front_heading("ACKNOWLEDGEMENT", "bm_ack", break_before=True, after=18)
para(
    "I would like to express my sincere gratitude to our seminar guide, Miss. Sakshi Ambadkar, "
    "for her continuous guidance, patient review and constructive criticism throughout the "
    "preparation of this report. Her questions repeatedly pushed us to separate what the "
    "prototype actually demonstrates from what we hoped it would demonstrate.",
    after=10,
)
para(
    "I am thankful to the faculty members and the Department of Computer Science and Engineering "
    "at NMIET for providing the infrastructure, laboratory access and academic freedom needed to "
    "design, build and deploy this project. I also acknowledge the open-source communities behind "
    "React, TypeScript, Vite, Tailwind CSS, Zustand and the Python libraries used to prepare this "
    "document; their documentation made a small, independent implementation possible.",
    after=10,
)
para(
    "Finally, I extend my heartfelt thanks to my family, teammates and peers for their constant "
    "motivation, honest feedback and support during the completion of this seminar. The working "
    "prototype, the deployment and the demonstration dataset were built collaboratively, and the "
    "credit for anything that works belongs to the team.",
    after=28,
)
para("Vivek Rathod (B49)", 11, WD_ALIGN_PARAGRAPH.RIGHT, after=3, line=1.2)
para("(T.Y. B.E., NMIET)", 11, WD_ALIGN_PARAGRAPH.RIGHT, after=0, line=1.2, italic=True)

# abstract
front_heading("ABSTRACT", "bm_abs", break_before=True, after=18)
para(
    "Last-mile shrinkage is usually investigated at the warehouse, after the vehicle has returned "
    "and the count has failed. This report argues that the failure begins earlier, at dispatch, "
    "when a buffer quantity leaves the warehouse without an enforceable relationship to the seal, "
    "the driver, the route and the evidence of what happened in transit. The result is a custody "
    "gap: the final count exists but has no authority, because the conditions for judging it were "
    "never recorded.",
    after=8,
)
para(
    "The report presents ChainGuard v2, a browser-delivered control plane that makes the custody "
    "of a last-mile consignment explicit from dispatch to warehouse return. Each consignment is "
    "bound to a unique digital seal and a declared buffer quantity. The seal can be broken only "
    "through a customer-issued single-use unlock token, and the exact buffer count must return. "
    "Driver telemetry, GPS-anchored incidents, the customer attestation and the warehouse count "
    "share one dispatch record, so reconciliation compares fields within a record rather than "
    "reconstructing four separate systems.",
    after=8,
)
para(
    "The prototype implements four role portals — vendor dispatch, driver telemetry, customer "
    "verification and audit reconciliation — over a single reactive Zustand store. It formalises "
    "custody as a guarded state machine, computes a unit-exact variance, raises a priority-1 "
    "shrinkage alert and locks the driver’s settlement on any non-zero variance, and derives fleet "
    "KPIs and a driver trust score from the same records. A seeded fleet of six dispatches covers "
    "every lifecycle state, and a reproducible fraud scenario demonstrates the complete path from "
    "seal issuance to payout lock.",
    after=8,
)
para(
    "The report also states the prototype’s limits. The trust boundary is the browser, identifiers "
    "are not cryptographic secrets, settlement is a flag rather than a payment instruction, and the "
    "audit log is not yet immutable. These limitations define a concrete future scope: a server-side "
    "custody service, signed tokens and hardware seals, a hash-chained evidence log and integration "
    "with settlement systems. The broader contribution is a reusable custody pattern, not a product: "
    "one active seal, consent before break, buffer as expected return and variance as a settlement "
    "gate.",
    after=0,
)

# table of contents
front_heading("TABLE OF CONTENTS", "bm_toc", break_before=True, after=14)

front_entries = [
    ("Certificate", "bm_cert"),
    ("Acknowledgement", "bm_ack"),
    ("Abstract", "bm_abs"),
    ("Table of Contents", "bm_toc"),
    ("List of Figures", "bm_lof"),
    ("List of Tables", "bm_lot"),
    ("List of Abbreviations", "bm_abbr"),
]
for label, bm in front_entries:
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.05
    pf.tab_stops.add_tab_stop(
        Inches(TEXT_W), alignment=WD_TAB_ALIGNMENT.RIGHT, leader=WD_TAB_LEADER.DOTS
    )
    run = p.add_run(label)
    style_run(run, 11.3)
    add_field(p, f"PAGEREF {bm} \\h", "", size=11.3, tab_first=True)

spacer(3)
toc_p = doc.add_paragraph()
toc_p.paragraph_format.space_before = Pt(0)
toc_p.paragraph_format.space_after = Pt(0)
toc_p.paragraph_format.line_spacing = 1.0
add_field(
    toc_p,
    'TOC \\o "1-2" \\z \\u',
    "Update the field to build the table of contents.",
    size=11.3,
)

# figures and tables known before the body is written
figures = []
tables = []
for block in BLOCKS:
    if block[0] == "fig":
        label, body = caption_parts(block[2])
        figures.append((label, body, "fig_" + label.split()[-1].replace(".", "_")))
    elif block[0] == "tbl":
        label, body = caption_parts(block[1]["caption"])
        tables.append((label, body, "tbl_" + label.split()[-1].replace(".", "_")))

# list of figures and list of tables (single page, template style)
front_heading("LIST OF FIGURES", "bm_lof", break_before=True, size=15.5, after=14)


def list_table(rows, first_col):
    """Figure/Table list in the template's bordered, shaded-header style."""
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    cell_margins(table, left=0.08, right=0.08, top=0.0, bottom=0.0)

    widths = (1.10, 3.55, 1.05)
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)

    heading_cells = table.rows[0].cells
    for i, label in enumerate((first_col, "Name", "Page no")):
        cell = heading_cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade(cell, "DDEBF7")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(label)
        style_run(run, 13, bold=True)
    header_row = table.rows[0]._tr.get_or_add_trPr()
    header_row.append(OxmlElement("w:cantSplit"))
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), "500")
    height.set(qn("w:hRule"), "atLeast")
    header_row.append(height)

    for label, name, bm in rows:
        row = table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        height = OxmlElement("w:trHeight")
        height.set(qn("w:val"), "460")
        height.set(qn("w:hRule"), "atLeast")
        tr_pr.append(height)
        for i, value in enumerate((label, name)):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(value)
            style_run(run, 11)
        cell = row.cells[2]
        cell.width = Inches(widths[2])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        add_field(p, f"PAGEREF {bm} \\h", "", size=11)
    return table


list_table(figures, "Figure No")
spacer(24)
front_heading("LIST OF TABLES", "bm_lot", size=15.5, after=14)
list_table(tables, "Table No")

# abbreviations (own section so the running header appears above the list)
abbr_section = new_section()
add_running_header(abbr_section)
add_page_number_footer(abbr_section)
set_page_numbering(abbr_section, "upperRoman")
front_heading("LIST OF ABBREVIATIONS", "bm_abbr", after=14)
abbreviations = [
    ("API", "Application Programming Interface"),
    ("CDN", "Content Delivery Network"),
    ("CI/CD", "Continuous Integration / Continuous Delivery"),
    ("CSPRNG", "Cryptographically Secure Pseudo-Random Number Generator"),
    ("CSV", "Comma-Separated Values"),
    ("DOM", "Document Object Model"),
    ("DPDP", "Digital Personal Data Protection (Act, 2023)"),
    ("GPS", "Global Positioning System"),
    ("HMAC", "Keyed-Hash Message Authentication Code"),
    ("HTTP", "Hypertext Transfer Protocol"),
    ("INR", "Indian Rupee"),
    ("IoT", "Internet of Things"),
    ("JSON", "JavaScript Object Notation"),
    ("JWT", "JSON Web Token"),
    ("KPI", "Key Performance Indicator"),
    ("NIST", "National Institute of Standards and Technology"),
    ("P1", "Priority One"),
    ("POD", "Proof of Delivery"),
    ("QR", "Quick Response"),
    ("RBAC", "Role-Based Access Control"),
    ("REST", "Representational State Transfer"),
    ("SDK", "Software Development Kit"),
    ("SPA", "Single-Page Application"),
    ("SRS", "Software Requirements Specification"),
    ("TLS", "Transport Layer Security"),
    ("TMS", "Transport Management System"),
    ("UI", "User Interface"),
    ("URL", "Uniform Resource Locator"),
    ("UTC", "Coordinated Universal Time"),
    ("UUID", "Universally Unique Identifier"),
    ("W3C", "World Wide Web Consortium"),
    ("XSS", "Cross-Site Scripting"),
    ("ZTA", "Zero Trust Architecture"),
]
abbr = doc.add_table(rows=1, cols=2)
abbr.alignment = WD_TABLE_ALIGNMENT.CENTER
abbr.autofit = False

tbl_pr = abbr._tbl.tblPr
borders = OxmlElement("w:tblBorders")
for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "6")
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    borders.append(el)
tbl_pr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")

cell_margins(abbr, left=0.04, right=0.04, top=0.03, bottom=0.03)

for i, width in enumerate((1.05, 4.65)):
    abbr.columns[i].width = Inches(width)

# Shade header row and set header font
heading_cells = abbr.rows[0].cells
for i, label in enumerate(("Abbreviation", "Full Form")):
    cell = heading_cells[i]
    cell.width = Inches((1.05, 4.65)[i])
    shade(cell, "DDEBF7")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(label)
    style_run(run, 9.5, bold=True)

for short, full in abbreviations:
    row = abbr.add_row()
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, value in enumerate((short, full)):
        cell = row.cells[i]
        cell.width = Inches((1.05, 4.65)[i])
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(value)
        style_run(run, 9.9)


# ---------------------------------------------------------------------------
# body helpers
# ---------------------------------------------------------------------------
def add_caption(text, bookmark=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.space_before = Pt(4)
    pf.space_after = Pt(10)
    pf.line_spacing = 1.0
    run = p.add_run(text)
    style_run(run, size=CAPTION_SIZE, bold=True)
    if bookmark:
        add_bookmark(p, bookmark)
    return p


def set_cell(cell, text, size=TABLE_SIZE, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(1)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.0
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    add_rich(p, str(text), size=size, bold_all=bold)


def add_table(spec):
    headers = spec["headers"]
    rows = spec["rows"]
    widths = list(spec["widths"])
    size = spec.get("size", 8.8)
    if spec["caption"].startswith("Table 3.1"):
        widths = [1.0, 1.5, 1.2, 1.0, 1.05]
        size = 6.8

    scale = TEXT_W / sum(widths)
    widths = [w * scale for w in widths]

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    no_borders(table)
    cell_margins(table, left=0.04, right=0.04, top=0.03, bottom=0.03)
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)

    header = table.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))
    tr_pr.append(OxmlElement("w:cantSplit"))
    for i, label in enumerate(headers):
        cell = header.cells[i]
        cell.width = Inches(widths[i])
        set_cell(cell, label, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        bottom_rule(cell, size="4", color="000000")

    for values in rows:
        row = table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            set_cell(cell, value, size=size)

    caption = caption_text(spec["caption"])
    label, _ = caption_parts(caption)
    bm = "tbl_" + label.split()[-1].replace(".", "_")
    add_caption(caption, bm)


def add_figure(filename, caption):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.space_before = Pt(8)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(FIG / filename), width=Inches(5.0))

    caption = caption_text(caption)
    label, _ = caption_parts(caption)
    bm = "fig_" + label.split()[-1].replace(".", "_")
    add_caption(caption, bm)


def add_formula(lines):
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(3)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.15
        run = p.add_run(line)
        style_run(run, size=11.3, italic=True)


def add_code(lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(TEXT_W)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge in ("top", "bottom", "left", "right"):
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "BFBFBF")
        else:
            el.set(qn("w:val"), "nil")
        el.set(qn("w:space"), "0")
        borders.append(el)
    tbl_pr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    cell_margins(table, left=0.1, right=0.1, top=0.06, bottom=0.06)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    cell = table.rows[0].cells[0]
    cell.width = Inches(TEXT_W)
    shade(cell, "F5F5F5")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
    for line in lines:
        p = cell.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(line if line else " ")
        style_run(run, size=8.8, name=MONO_FONT)
    spacer(5)


def add_callout(text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(TEXT_W)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge == "left":
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "18")
            el.set(qn("w:color"), "404040")
        else:
            el.set(qn("w:val"), "nil")
        el.set(qn("w:space"), "0")
        borders.append(el)
    tbl_pr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    cell_margins(table, left=0.14, right=0.12, top=0.07, bottom=0.07)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    cell = table.rows[0].cells[0]
    cell.width = Inches(TEXT_W)
    shade(cell, "F2F2F2")
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before = Pt(1)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.2
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_rich(p, text, size=10.5)
    spacer(8)


def chapter_section(number, title):
    section = new_section()
    section.different_first_page_header_footer = True
    add_running_header(section, f"Chapter {number}", first_page=True)
    add_running_header(section)
    add_page_number_footer(section, first_page=True)
    add_page_number_footer(section)
    if number == 1:
        set_page_numbering(section, "decimal", 1)
    else:
        set_page_numbering(section, "decimal", None)
    p = doc.add_heading(title, level=1)
    p.paragraph_format.page_break_before = False
    return section


# ---------------------------------------------------------------------------
# chapters
# ---------------------------------------------------------------------------
chapter_no = 0
ref_no = 0
for block in BLOCKS:
    kind = block[0]

    if kind == "h1":
        raw = block[1].strip()
        number = raw.split(".")[0].strip()
        if raw.upper().startswith("10."):
            section = new_section()
            section.different_first_page_header_footer = False
            add_running_header(section)
            add_page_number_footer(section)
            set_page_numbering(section, "decimal", None)
            p = doc.add_heading("References", level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(18)
            for run in p.runs:
                style_run(run, size=15, bold=True, caps=True)
        else:
            chapter_no += 1
            title = CHAPTER_TITLES.get(number, raw.title())
            chapter_section(chapter_no, title)
    elif kind == "h2":
        doc.add_heading(block[1], level=2)
    elif kind == "h3":
        doc.add_heading(block[1], level=3)
    elif kind == "p":
        para(block[1])
    elif kind == "bul":
        for item in block[1]:
            add_item("•", item)
        spacer(3)
    elif kind == "num":
        for index, item in enumerate(block[1], start=1):
            add_item(f"{index}.", item, left=0.4, hang=0.28)
        spacer(3)
    elif kind == "fig":
        add_figure(block[1], block[2])
    elif kind == "tbl":
        add_table(block[1])
    elif kind == "code":
        add_code(block[1])
    elif kind == "callout":
        add_callout(block[1])
    elif kind == "formula":
        add_formula(block[1])
    elif kind == "ref":
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.4)
        pf.first_line_indent = Inches(-0.4)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(0)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.05
        pf.tab_stops.add_tab_stop(Inches(0.4))
        ref_no += 1
        run = p.add_run(f"[{ref_no}]\t")
        style_run(run, 10.5, bold=True)
        add_rich(p, block[1], size=10.5)

doc.save(OUT_DOCX)
print("wrote", OUT_DOCX)
print("blocks:", len(BLOCKS), "sections:", len(doc.sections))
print("figures:", len(figures), "tables:", len(tables))
